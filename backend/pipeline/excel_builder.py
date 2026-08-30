"""Gera o workbook no layout provisório. Ver docs/PJE_CALC_LAYOUT.md."""

from __future__ import annotations

from io import BytesIO

import pendulum
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

from schemas.extraction import ExtractionResult

LAYOUT_VERSION = "provisional-0.1"

CARTAO_COLUMNS = [
    "competencia",
    "data",
    "entrada",
    "saida",
    "intervalo_inicio",
    "intervalo_fim",
    "origem_pagina",
    "confianca",
]

HOLERITE_COLUMNS = [
    "competencia",
    "verba",
    "valor",
    "origem_pagina",
    "confianca",
]

HEADER_FILL = PatternFill("solid", fgColor="1B3A4B")
HEADER_FONT = Font(color="FFFFFF", bold=True)


class ExcelBuilder:
    def build(self, extraction: ExtractionResult) -> bytes:
        raise NotImplementedError


class ProvisionalExcelBuilder(ExcelBuilder):
    """Layout hipotético. Não é o modelo oficial do PJe-Calc."""

    def build(self, extraction: ExtractionResult) -> bytes:
        workbook = Workbook()
        self._write_time_cards(workbook, extraction)
        self._write_payslips(workbook, extraction)
        self._write_session_meta(workbook, extraction)

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _write_time_cards(self, workbook: Workbook, extraction: ExtractionResult) -> None:
        sheet = workbook.active
        sheet.title = "CartaoPonto"
        self._header(sheet, CARTAO_COLUMNS)
        sheet["A1"].comment = Comment(
            f"HIPÓTESE {LAYOUT_VERSION} — não é o layout oficial do PJe-Calc",
            "tcc-pjecalc",
        )
        for row_index, entry in enumerate(extraction.time_cards, start=2):
            sheet.cell(row_index, 1, entry.competence)
            sheet.cell(row_index, 2, entry.date)
            sheet.cell(row_index, 3, entry.clock_in)
            sheet.cell(row_index, 4, entry.clock_out)
            sheet.cell(row_index, 5, entry.break_start)
            sheet.cell(row_index, 6, entry.break_end)
            sheet.cell(row_index, 7, entry.source_page)
            sheet.cell(row_index, 8, entry.confidence)
        self._autosize(sheet, len(CARTAO_COLUMNS))

    def _write_payslips(self, workbook: Workbook, extraction: ExtractionResult) -> None:
        sheet = workbook.create_sheet("Holerite")
        self._header(sheet, HOLERITE_COLUMNS)
        for row_index, entry in enumerate(extraction.payslips, start=2):
            sheet.cell(row_index, 1, entry.competence)
            sheet.cell(row_index, 2, entry.item_name)
            amount = None if entry.amount is None else round(float(entry.amount), 2)
            sheet.cell(row_index, 3, amount)
            sheet.cell(row_index, 4, entry.source_page)
            sheet.cell(row_index, 5, entry.confidence)
        self._autosize(sheet, len(HOLERITE_COLUMNS))

    def _write_session_meta(self, workbook: Workbook, extraction: ExtractionResult) -> None:
        sheet = workbook.create_sheet("MetadadosSessao")
        self._header(sheet, ["chave", "valor"])
        rows = [
            ("layout_version", LAYOUT_VERSION),
            ("schema_version", extraction.schema_version),
            ("generated_at", pendulum.now("UTC").to_iso8601_string()),
            ("time_card_rows", len(extraction.time_cards)),
            ("payslip_rows", len(extraction.payslips)),
            ("pdf_page_count", extraction.pdf_page_count),
            ("candidate_page_count", extraction.candidate_page_count),
            ("gemini_call_count", extraction.gemini_call_count),
        ]
        for row_index, (key, value) in enumerate(rows, start=2):
            sheet.cell(row_index, 1, key)
            sheet.cell(row_index, 2, value)
        self._autosize(sheet, 2)

    @staticmethod
    def _header(sheet, columns: list[str]) -> None:
        for index, name in enumerate(columns, start=1):
            cell = sheet.cell(1, index, name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

    @staticmethod
    def _autosize(sheet, column_count: int) -> None:
        for index in range(1, column_count + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 18
