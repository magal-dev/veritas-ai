from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

import fitz
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app
from pipeline.extractor import LocalDocumentExtractor
from pipeline.validator import SchemaExtractionValidator
from schemas.extraction import Conflict, ExtractionResult, TimeCardEntry
from schemas.jobs import JobStatus
from services.session_store import JobSession, session_store


def _pdf_bytes(pages: int = 3) -> bytes:
    document = fitz.open()
    for index in range(pages):
        page = document.new_page()
        page.insert_text((72, 72), f"Pagina {index + 1}")
    payload = document.tobytes()
    document.close()
    return payload


def _seed_session(extraction: ExtractionResult) -> None:
    session = JobSession(
        job_id=extraction.job_id,
        status=JobStatus.COMPLETED,
        created_at=datetime.now(timezone.utc),
        extraction=extraction,
    )
    session_store.put(session)


def test_health_reports_ok():
    client = TestClient(app)
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_job_flow_counts_pages_and_returns_empty_extraction(tmp_path):
    pdf_path = tmp_path / "processo.pdf"
    pdf_path.write_bytes(_pdf_bytes(4))
    signals = LocalDocumentExtractor().extract(pdf_path)
    assert signals.pdf_page_count == 4
    assert signals.candidate_pages == []

    client = TestClient(app)
    response = client.post(
        "/api/v1/jobs",
        files={"file": ("processo.pdf", BytesIO(_pdf_bytes(4)), "application/pdf")},
    )
    assert response.status_code == 200, response.text
    job_id = response.json()["job_id"]
    assert response.json()["pdf_page_count"] == 4

    preview = client.get(f"/api/v1/jobs/{job_id}/preview")
    assert preview.status_code == 200
    extraction = preview.json()["extraction"]
    assert extraction["time_cards"] == []
    assert extraction["payslips"] == []
    assert extraction["gemini_call_count"] == 0

    excel = client.get(f"/api/v1/jobs/{job_id}/excel")
    assert excel.status_code == 200
    assert excel.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    discarded = client.delete(f"/api/v1/jobs/{job_id}")
    assert discarded.status_code == 204
    missing = client.get(f"/api/v1/jobs/{job_id}")
    assert missing.status_code == 404


def test_rejects_non_pdf():
    client = TestClient(app)
    response = client.post(
        "/api/v1/jobs",
        files={"file": ("notas.txt", b"nao e pdf", "text/plain")},
    )
    assert response.status_code == 400
    assert response.json()["detail"]["error_code"] == "INVALID_FILE_TYPE"


def test_empty_result_schema_roundtrip():
    result = ExtractionResult(job_id=uuid4())
    validated = SchemaExtractionValidator().validate(result)
    assert validated.schema_version == "extraction-schema-0.1"


def test_patch_extraction_normalizes_time():
    job_id = uuid4()
    extraction = ExtractionResult(
        job_id=job_id,
        time_cards=[
            TimeCardEntry(
                date="2024-03-04",
                clock_in="8h00",
                clock_out="17:00",
                source_page=1,
                confidence=0.8,
            )
        ],
    )
    _seed_session(extraction)

    client = TestClient(app)
    response = client.patch(
        f"/api/v1/jobs/{job_id}/extraction",
        json={
            "time_cards": [
                {
                    "date": "2024-03-04",
                    "competence": None,
                    "clock_in": "8h00",
                    "clock_out": "17:00",
                    "break_start": None,
                    "break_end": None,
                    "source_page": 1,
                    "confidence": 0.8,
                    "missing_fields": [],
                    "ambiguous_fields": [],
                }
            ],
            "payslips": [],
        },
    )
    assert response.status_code == 200, response.text
    card = response.json()["extraction"]["time_cards"][0]
    assert card["clock_in"] == "08:00"


def test_patch_resolves_conflict_and_excel_has_single_row():
    job_id = uuid4()
    extraction = ExtractionResult(
        job_id=job_id,
        time_cards=[
            TimeCardEntry(
                date="2024-03-04",
                clock_in="08:00",
                clock_out="17:00",
                source_page=1,
                confidence=0.9,
            ),
            TimeCardEntry(
                date="2024-03-04",
                clock_in="08:00",
                clock_out="18:00",
                source_page=2,
                confidence=0.9,
            ),
        ],
        conflicts=[
            Conflict(
                field="time_cards[2024-03-04]",
                values=[
                    {
                        "clock_in": "08:00",
                        "clock_out": "17:00",
                        "break_start": None,
                        "break_end": None,
                    },
                    {
                        "clock_in": "08:00",
                        "clock_out": "18:00",
                        "break_start": None,
                        "break_end": None,
                    },
                ],
                source_pages=[1, 2],
                note="Horários diferentes para o dia 2024-03-04 em páginas distintas.",
            )
        ],
    )
    _seed_session(extraction)

    client = TestClient(app)
    blocked = client.get(f"/api/v1/jobs/{job_id}/excel")
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["error_code"] == "CONFLICTS_UNRESOLVED"

    response = client.patch(
        f"/api/v1/jobs/{job_id}/extraction",
        json={
            "time_cards": [
                {
                    "date": "2024-03-04",
                    "competence": None,
                    "clock_in": "08:00",
                    "clock_out": "17:00",
                    "break_start": None,
                    "break_end": None,
                    "source_page": 1,
                    "confidence": 0.9,
                    "missing_fields": [],
                    "ambiguous_fields": [],
                }
            ],
            "payslips": [],
        },
    )
    assert response.status_code == 200
    assert response.json()["extraction"]["conflicts"] == []

    excel = client.get(f"/api/v1/jobs/{job_id}/excel")
    assert excel.status_code == 200
    workbook = load_workbook(BytesIO(excel.content))
    sheet = workbook["CartaoPonto"]
    data_rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[1]]
    assert len(data_rows) == 1


def test_patch_unknown_session_returns_404():
    client = TestClient(app)
    response = client.patch(
        f"/api/v1/jobs/{uuid4()}/extraction",
        json={"time_cards": [], "payslips": []},
    )
    assert response.status_code == 404
    assert response.json()["detail"]["error_code"] == "SESSION_NOT_FOUND"
