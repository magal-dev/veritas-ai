from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from pipeline.excel_builder import LAYOUT_VERSION, ProvisionalExcelBuilder
from schemas.extraction import ExtractionResult, PayslipEntry, TimeCardEntry
from pipeline.validator import SchemaExtractionValidator


def test_provisional_workbook_has_expected_sheets_and_version():
    extraction = ExtractionResult(job_id=uuid4(), pdf_page_count=12)
    payload = ProvisionalExcelBuilder().build(extraction)
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == ["CartaoPonto", "Holerite", "MetadadosSessao"]
    meta = {row[0]: row[1] for row in workbook["MetadadosSessao"].iter_rows(min_row=2, values_only=True)}
    assert meta["layout_version"] == LAYOUT_VERSION
    assert meta["pdf_page_count"] == 12
    assert meta["time_card_rows"] == 0
    assert meta["gemini_call_count"] == 0


def test_validator_accepts_empty_result():
    result = ExtractionResult(job_id=uuid4())
    validated = SchemaExtractionValidator().validate(result)
    assert validated.time_cards == []
    assert validated.payslips == []


def test_excel_rows_follow_provisional_columns():
    extraction = ExtractionResult(
        job_id=uuid4(),
        pdf_page_count=3,
        time_cards=[
            TimeCardEntry(
                date="2024-03-01",
                competence="2024-03",
                clock_in="08:00",
                clock_out="17:00",
                break_start="12:00",
                break_end="13:00",
                source_page=2,
                confidence=0.8,
            )
        ],
        payslips=[
            PayslipEntry(
                competence="2024-03",
                item_name="Salário-base",
                amount=2500.5,
                source_page=3,
                confidence=0.7,
            )
        ],
    )
    payload = ProvisionalExcelBuilder().build(extraction)
    workbook = load_workbook(BytesIO(payload))
    ponto = next(workbook["CartaoPonto"].iter_rows(min_row=2, values_only=True))
    assert ponto[1] == "2024-03-01"
    assert ponto[2] == "08:00"
    holerite = next(workbook["Holerite"].iter_rows(min_row=2, values_only=True))
    assert holerite[1] == "Salário-base"
    assert holerite[2] == 2500.5
